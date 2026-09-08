# -*- coding: utf-8 -*-
"""生成《周五课堂系统架构设计图》Word 文档（python-docx）。"""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))


def set_font(run, ascii_name='Times New Roman', east_name='宋体', size=10.5, bold=False):
    run.font.name = ascii_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), east_name)
    run.font.size = Pt(size)
    run.bold = bold


doc = Document()

# 正文默认字体：宋体 10.5pt
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

# 标题
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('周五课堂系统架构设计图')
set_font(run, east_name='黑体', size=16, bold=True)

# 副标题
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub.add_run('智能教学互动平台')
set_font(run, east_name='黑体', size=12)

# 图注
cap = doc.add_paragraph()
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = cap.add_run('图 1  周五课堂系统架构图')
set_font(run, east_name='黑体', size=10.5)

# 插入架构图
img_path = os.path.join(BASE, 'architecture.png')
doc.add_picture(img_path, width=Cm(16))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

# 分层说明
h = doc.add_paragraph()
run = h.add_run('分层说明')
set_font(run, east_name='黑体', size=14, bold=True)

layers = [
    ('客户端层（前端 Vue3 + Vite · JavaScript）',
     '教师端与学生端，分别负责课件上传、直播授课、翻页控制，以及直播观看、AI 文字问答交互。'),
    ('应用层（后端 SpringBoot 3.5）',
     'Controller 提供 REST API；WebSocket 负责翻页页码广播；Service 承载 F001~F006 业务逻辑；'
     'Repository 通过 Spring Data JPA 访问数据库。'),
    ('数据层',
     'MySQL 9.5 存储结构化数据（friday_class 库，9 张表）；文件系统存储 .pptx 课件与网页幻灯片。'),
    ('外部服务',
     'DeepSeek API 提供纯文本大模型能力，用于 AI 课件解析、学生问答与课后总结生成。'),
]
for name, text in layers:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(name + '：')
    set_font(run, east_name='黑体', size=10.5, bold=True)
    run2 = p.add_run(text)
    set_font(run2, east_name='宋体', size=10.5)

out = os.path.join(BASE, '周五课堂_架构设计图.docx')
doc.save(out)
print('已生成:', out)
