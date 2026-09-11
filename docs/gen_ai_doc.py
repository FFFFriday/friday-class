# -*- coding: utf-8 -*-
"""生成《周五课堂_AI智能体设计》Word 文档（python-docx）。

字体规范沿用项目既有交付件：标题黑体、正文宋体 10.5pt。
运行前请先用 gen_diagram.py 渲染出 docs 下的 PNG。
"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = os.path.dirname(os.path.abspath(__file__))


def set_font(run, east_name="宋体", ascii_name="Times New Roman", size=10.5, bold=False):
    run.font.name = ascii_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_name)
    run.font.size = Pt(size)
    run.bold = bold


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    size = 14 if level == 1 else 12
    run = p.add_run(text)
    set_font(run, east_name="黑体", size=size, bold=True)
    return p


def body(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, east_name="宋体", size=10.5)
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.space_after = Pt(6)
    return p


def bullet(doc, name, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(name + "：")
    set_font(run, east_name="黑体", size=10.5, bold=True)
    run2 = p.add_run(text)
    set_font(run2, east_name="宋体", size=10.5)
    return p


def picture(doc, filename, caption, width_cm):
    path = os.path.join(BASE, filename)
    doc.add_picture(path, width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    set_font(run, east_name="黑体", size=10.5)
    return cap


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_font(run, east_name="黑体", size=10.5, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(val)
            set_font(run, east_name="宋体", size=10.5)
    return t


doc = Document()

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

# 标题
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(title.add_run("周五课堂 AI 对话智能体设计"), east_name="黑体", size=16, bold=True)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(sub.add_run("智能教学互动平台"), east_name="黑体", size=12)

# 一、概述
heading(doc, "一、总体设计")
body(doc, "平台对人工智能的调用集中在两处，二者共用同一套「提示词组装 — 模型调用 — 失败重试 — 结果落库」"
          "骨架，因此统一抽象为「智能体」。二者的区别只在于角色与输入输出：")
table(doc, ["智能体", "对应功能", "运行时机", "输入", "输出"], [
    ["课件解析智能体", "F002", "教师上传课件后（异步）", "某一页的文字", "知识点 + 预置提问"],
    ["学生问答智能体", "F004", "学生提问时（实时）", "本页知识点 + 预置提问 + 学生问题", "一段回答"],
])
body(doc, "一句话概括：解析智能体把「课件」变成「结构化知识」，问答智能体再拿这份知识回答学生。"
          "所采用的大模型为 DeepSeek 纯文本模型，不涉及音频、图像与多模态能力。")

# 二、图
heading(doc, "二、设计图")
body(doc, "图 1 描述学生端一次完整问答的全过程，包含课前提示词预下发、课中翻页切换上下文、"
          "以及提问成功与失败两条分支。")
picture(doc, "ai-qa-sequence.png", "图 1  学生端 AI 问答时序图（F004）", 14)

doc.add_page_break()
body(doc, "图 2 描述课件解析的全过程。解析是异步的：教师上传后立即返回，后台逐页调用模型，"
          "因此一份几十页的课件不会阻塞上传接口。")
picture(doc, "ai-parse-sequence.png", "图 2  课件解析时序图（F001 + F002，异步）", 14)

doc.add_page_break()
body(doc, "图 3 描述问答智能体在调用模型之前，如何一步步组装提示词，以及每一步的降级出口。")
picture(doc, "ai-prompt-flow.png", "图 3  AI 问答智能体的提示词组装流程", 13)

# 三、核心机制
doc.add_page_break()
heading(doc, "三、核心机制：靠页码定位上下文")
body(doc, "本设计最关键的一环，是把「老师正在讲第几页」作为上下文定位的依据：")
bullet(doc, "翻页广播", "老师端把课件转成网页幻灯片，翻页时通过 WebSocket 广播当前页码。")
bullet(doc, "页码定位", "学生端收到页码后，切换到该页对应的知识点上下文。")
bullet(doc, "带页提问", "学生提问时携带当前页的标识，后端只取该页知识点组装提示词。")
body(doc, "这样做的收益是：问答智能体不需要「理解整份课件」，只需聚焦当前这一页。上下文小、响应快、"
          "成本低，而且回答天然贴合老师正在讲的内容，不会答到别处去。")

heading(doc, "四、提示词设计要点")
table(doc, ["设计点", "做法", "原因"], [
    ["角色约束", "声明为「课堂助教」，只回答当前页相关问题", "收窄回答范围，避免跑题"],
    ["输入隔离", "课件正文与学生问题分别用标记包裹", "降低正文或问题中夹带指令造成的干扰"],
    ["结构化输出", "解析智能体强制返回 JSON", "便于程序解析，解析失败可自动重试"],
    ["允许空结果", "封面、目录页返回空数组", "不为凑数而编造知识点"],
    ["输出长度", "回答限制在 200 字以内", "课堂上长篇回答无意义，也控制成本"],
    ["前端渲染", "以纯文本插值展示，不使用富文本注入", "避免模型输出被当作 HTML 执行"],
])

heading(doc, "五、超时与重试预算")
body(doc, "在线问答与离线解析的场景不同，采用两套预算，不共用参数：")
table(doc, ["场景", "总预算", "单次超时", "重试次数"], [
    ["问答（在线）", "不超过 20 秒", "8 秒", "最多 1 次（退避 1 秒）"],
    ["解析（离线）", "无硬上限", "60 秒", "最多 3 次（1、2、4 秒）"],
])
body(doc, "此外，仅对超时、限流（429）与服务端错误（5xx）重试；"
          "鉴权失败（401）、参数错误（400）、余额不足（402）属于不可恢复错误，立即失败并记录错误码。")

heading(doc, "六、异常与降级")
table(doc, ["场景", "处理方式"], [
    ["单页解析失败", "指数退避重试；重试后仍失败则记录错误、跳过该页继续下一页"],
    ["部分页失败", "任务标记为「部分成功」，课件仍标记为已解析，并置任务级标记提示教师"],
    ["全部页失败", "任务与课件均标记为失败状态，供教师触发重试"],
    ["课件正在解析中", "提示「课件还在解析中，稍后再试」，不调用模型"],
    ["课件解析失败", "提示「本页解析失败，请告诉老师」，不调用模型"],
    ["本页确无知识点", "提示「本页暂无解析内容」，不调用模型"],
    ["问答调用超时或失败", "提示「AI 助教暂时忙不过来，请稍后再试」，并写入一条失败状态的问答记录"],
])

heading(doc, "七、并发与限流")
table(doc, ["措施", "说明"], [
    ["每学生令牌桶", "例如每 5 秒最多提问 1 次，超限时友好拒绝"],
    ["有界线程池", "限制同时进行的模型调用数，超出的请求排队"],
    ["排队上限", "队列满时快速失败并提示，不无限堆积"],
    ["限流退避", "收到限流响应时指数退避重试，不直接判为失败"],
])

heading(doc, "八、成本控制")
table(doc, ["手段", "效果"], [
    ["解析结果落库", "同一份课件只解析一次，之后复用，不重复消耗额度"],
    ["上下文只带当前页", "相比把整份课件塞进提示词，消耗降低一个数量级"],
    ["相同问题去重", "同一问题多人提问只调用一次模型，课堂场景收益最大"],
    ["解析多页合并调用", "摊薄重复的系统提示词开销"],
    ["无知识点时不调用", "空白页直接返回，省下一次调用"],
    ["逐页增量解析", "单页失败无需重跑整份课件"],
])

out = os.path.join(BASE, "周五课堂_AI智能体设计.docx")
doc.save(out)
print("已生成:", out)
